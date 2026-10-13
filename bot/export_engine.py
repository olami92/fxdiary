import os
import datetime
import logging
import re
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from typing import List, Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

logger = logging.getLogger(__name__)

def generate_pdf_report(user_info: Dict[str, Any], trades: List[Any], stats: Dict[str, Any], output_path: str) -> bool:
    """Generate a PDF Trading Journal Report."""
    try:
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        story = []
        styles = getSampleStyleSheet()

        # Custom Styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'DocSubTitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=15
        )
        section_style = ParagraphStyle(
            'SectionTitle',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=14,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=15,
            spaceAfter=8
        )
        cell_style = ParagraphStyle(
            'CellText',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#334155'),
            alignment=1 # Center
        )
        header_cell_style = ParagraphStyle(
            'HeaderCellText',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=colors.white,
            alignment=1
        )

        # Header Title
        trader_name = user_info.get("first_name") or user_info.get("username") or "Trader"
        date_str = datetime.datetime.now().strftime("%B %d, %Y")
        story.append(Paragraph(f"FXDiary Trading Performance Report", title_style))
        story.append(Paragraph(f"Account: <b>{trader_name}</b> | Generated on: {date_str} | Total Trades: {stats['total_trades']}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=15))

        # 4 Key Metrics Cards Table
        pnl_val = stats['total_net_pnl']
        pnl_color = "#10B981" if pnl_val >= 0 else "#EF4444"
        pnl_sign = "+" if pnl_val >= 0 else ""

        cards_data = [
            [
                Paragraph("<font color='#64748B' size=8>NET PROFIT/LOSS</font><br/><font size=14 color='%s'><b>%s$%s</b></font>" % (pnl_color, pnl_sign, f"{abs(pnl_val):,.2f}"), styles['Normal']),
                Paragraph("<font color='#64748B' size=8>WIN RATE</font><br/><font size=14 color='#3B82F6'><b>%s%%</b></font>" % stats['win_rate'], styles['Normal']),
                Paragraph("<font color='#64748B' size=8>PROFIT FACTOR</font><br/><font size=14 color='#8B5CF6'><b>%s</b></font>" % stats['profit_factor'], styles['Normal']),
                Paragraph("<font color='#64748B' size=8>TOTAL TRADES</font><br/><font size=14 color='#0F172A'><b>%s</b></font>" % stats['total_trades'], styles['Normal']),
            ]
        ]
        cards_table = Table(cards_data, colWidths=[130, 130, 130, 130])
        cards_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
            ('INNERGRID', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(cards_table)
        story.append(Spacer(1, 15))

        # Secondary Stats Table
        sec_stats = [
            ["Wins", str(stats['wins']), "Avg Win", f"${stats['avg_win']:,.2f}", "Best Trade", f"${stats['best_trade']:,.2f}"],
            ["Losses", str(stats['losses']), "Avg Loss", f"${stats['avg_loss']:,.2f}", "Worst Trade", f"${stats['worst_trade']:,.2f}"],
            ["Break Even", str(stats['breakeven']), "Max Drawdown", f"${stats['max_drawdown']:,.2f}", "Total Pips", f"{stats['total_pips']} pips"]
        ]
        sec_table = Table(sec_stats, colWidths=[90, 80, 90, 80, 90, 90])
        sec_table.setStyle(TableStyle([
            ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#334155')),
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTNAME', (2,0), (2,-1), 'Helvetica-Bold'),
            ('FONTNAME', (4,0), (4,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#F1F5F9')),
        ]))
        story.append(sec_table)
        story.append(Spacer(1, 15))

        # Trade History Table Title
        story.append(Paragraph("Trade Log History", section_style))

        # Trade History Table Data
        table_headers = ["Date & Time", "Ticket", "Symbol", "Type", "Lots", "Open", "Close", "Pips", "Net PnL"]
        table_rows = [[Paragraph(h, header_cell_style) for h in table_headers]]

        for t in trades[:100]: # Limit top 100 in PDF
            close_time = t.close_time.strftime("%Y-%m-%d %H:%M") if hasattr(t, "close_time") and t.close_time else "-"
            ticket = t.ticket if hasattr(t, "ticket") and t.ticket else "-"
            sym = t.symbol.upper() if hasattr(t, "symbol") else t.get("symbol", "-")
            ttype = t.trade_type.upper() if hasattr(t, "trade_type") else t.get("trade_type", "-")
            lots = f"{t.volume:.2f}" if hasattr(t, "volume") else f"{t.get('volume', 0.01):.2f}"
            open_p = f"{t.open_price:.5f}" if hasattr(t, "open_price") else f"{t.get('open_price', 0):.5f}"
            close_p = f"{t.close_price:.5f}" if hasattr(t, "close_price") else f"{t.get('close_price', 0):.5f}"
            pips = f"{t.pips:.1f}" if hasattr(t, "pips") else f"{t.get('pips', 0):.1f}"
            pnl = t.net_profit if hasattr(t, "net_profit") else t.get("net_profit", 0.0)

            pnl_txt_color = "#15803D" if pnl > 0 else ("#B91C1C" if pnl < 0 else "#334155")
            pnl_formatted = f"<font color='{pnl_txt_color}'><b>${pnl:+,.2f}</b></font>"

            table_rows.append([
                Paragraph(close_time, cell_style),
                Paragraph(ticket, cell_style),
                Paragraph(sym, cell_style),
                Paragraph(ttype, cell_style),
                Paragraph(lots, cell_style),
                Paragraph(open_p, cell_style),
                Paragraph(close_p, cell_style),
                Paragraph(pips, cell_style),
                Paragraph(pnl_formatted, cell_style),
            ])

        history_table = Table(table_rows, colWidths=[90, 55, 55, 45, 40, 60, 60, 45, 70])
        history_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))

        story.append(history_table)
        doc.build(story)
        return True
    except Exception as e:
        print(f"PDF Generation Error: {e}")
        return False

def generate_excel_report(user_info: Dict[str, Any], trades: List[Any], stats: Dict[str, Any], output_path: str) -> bool:
    """Create a styled MyFXkeeper report with period analytics and full trade history."""
    try:
        def value(trade, key, default=None):
            return getattr(trade, key, default) if not isinstance(trade, dict) else trade.get(key, default)

        def as_utc_naive(value_):
            if not isinstance(value_, datetime.datetime):
                return None
            if value_.tzinfo is not None:
                return value_.astimezone(datetime.timezone.utc).replace(tzinfo=None)
            return value_

        try:
            report_timezone = ZoneInfo(user_info.get("timezone") or "UTC")
        except (ZoneInfoNotFoundError, ValueError):
            report_timezone = datetime.timezone.utc
        now_local = datetime.datetime.now(report_timezone)
        now = now_local.astimezone(datetime.timezone.utc).replace(tzinfo=None)
        today = now_local.date()
        def local_midnight(day):
            local_time = datetime.datetime.combine(day, datetime.time.min, tzinfo=report_timezone)
            return local_time.astimezone(datetime.timezone.utc).replace(tzinfo=None)
        starts = {
            "WTD": local_midnight(today - datetime.timedelta(days=today.weekday())),
            "MTD": local_midnight(today.replace(day=1)),
            "YTD": local_midnight(today.replace(month=1, day=1)),
        }
        closed_trades = [t for t in trades if str(value(t, "status", "CLOSED") or "CLOSED").upper() != "OPEN"]

        def period_metrics(start):
            period_trades = [t for t in closed_trades if (as_utc_naive(value(t, "close_time")) or datetime.datetime.min) >= start and (as_utc_naive(value(t, "close_time")) or datetime.datetime.min) <= now]
            pnls = [float(value(t, "net_profit", 0.0) or 0.0) for t in period_trades]
            wins = sum(1 for pnl in pnls if pnl > 0.001)
            losses = sum(1 for pnl in pnls if pnl < -0.001)
            be = len(pnls) - wins - losses
            return {"pnl": sum(pnls), "trades": len(pnls), "wins": wins, "losses": losses, "breakeven": be,
                    "win_rate": (wins / len(pnls) * 100) if pnls else 0.0}

        periods = {name: period_metrics(start) for name, start in starts.items()}
        currency_codes = {"USD", "EUR", "GBP", "JPY", "CHF", "CAD", "AUD", "NZD", "CNY", "CNH", "HKD", "SGD", "SEK", "NOK", "MXN", "ZAR", "TRY", "PLN", "THB", "HUF"}
        currency_pnl: Dict[str, float] = {}
        pair_pnl: Dict[str, float] = {}
        for trade in closed_trades:
            symbol = str(value(trade, "symbol", "") or "").upper()
            normalized_symbol = re.sub(r"[^A-Z]", "", symbol)
            pnl = float(value(trade, "net_profit", 0.0) or 0.0)
            pair_pnl[symbol] = pair_pnl.get(symbol, 0.0) + pnl
            if len(normalized_symbol) >= 6:
                base, quote = normalized_symbol[:3], normalized_symbol[3:6]
                if base in currency_codes and quote in currency_codes:
                    direction = 1 if str(value(trade, "trade_type", "BUY")).upper() == "BUY" else -1
                    currency_pnl[base] = currency_pnl.get(base, 0.0) + pnl * direction
                    currency_pnl[quote] = currency_pnl.get(quote, 0.0) - pnl * direction

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Performance"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A5"
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        navy = "18201A"
        green = "22A249"
        pale_green = "E7F3E8"
        pale_red = "F7EBED"
        gray = "858A83"
        border_color = "E3E7E1"
        header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="252B25")
        thin_border = Border(bottom=Side(style="thin", color=border_color))
        money_format = '$#,##0.00;[Red]-$#,##0.00;–'
        percent_format = '0.0%'

        ws.merge_cells("A1:F1")
        ws["A1"] = "MyFXkeeper · Trading Performance Report"
        ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=navy)
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 38
        for cell in ws[1][:6]:
            cell.fill = PatternFill("solid", fgColor=navy)
        account_name = user_info.get("first_name") or user_info.get("username") or "Trader"
        ws.merge_cells("A2:F2")
        ws["A2"] = f"Account: {account_name}    ·    Generated: {now_local:%B %d, %Y %Z}"
        ws["A2"].font = Font(name="Aptos", size=10, color=gray, italic=True)

        def section(row, title, end_col=6):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
            cell = ws.cell(row=row, column=1, value=title)
            cell.font = Font(name="Aptos", size=12, bold=True, color=navy)
            cell.fill = PatternFill("solid", fgColor="E7F3E8")
            cell.alignment = Alignment(vertical="center")
            ws.row_dimensions[row].height = 24

        def header_row(row, labels):
            for col, label in enumerate(labels, start=1):
                cell = ws.cell(row=row, column=col, value=label)
                cell.font = header_font
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row].height = 23

        section(4, "PERIOD PERFORMANCE")
        header_row(5, ["Period", "Net P&L", "Trades", "Wins", "Losses", "Win rate"])
        for row, (name, label) in enumerate((("WTD", "Week to date"), ("MTD", "Month to date"), ("YTD", "Year to date")), start=6):
            metrics = periods[name]
            vals = [label, metrics["pnl"], metrics["trades"], metrics["wins"], metrics["losses"], metrics["win_rate"] / 100]
            for col, val in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = body_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left" if col == 1 else "right")
            ws.cell(row=row, column=2).number_format = money_format
            ws.cell(row=row, column=6).number_format = percent_format
            ws.cell(row=row, column=2).font = Font(name="Aptos", size=10, bold=True, color=green if metrics["pnl"] >= 0 else "9A3651")

        section(11, "OVERALL PERFORMANCE")
        header_row(12, ["Metric", "Value", "Metric", "Value", "Metric", "Value"])
        overall = [
            ("Net P&L", stats.get("total_net_pnl", 0), "Wins", stats.get("wins", 0), "Losses", stats.get("losses", 0)),
            ("Win rate", (stats.get("win_rate", 0) or 0) / 100, "Break even", stats.get("breakeven", 0), "Total trades", stats.get("total_trades", 0)),
            ("Profit factor", stats.get("profit_factor", 0), "Average win", stats.get("avg_win", 0), "Average loss", stats.get("avg_loss", 0)),
            ("Best trade", stats.get("best_trade", 0), "Worst trade", stats.get("worst_trade", 0), "Max drawdown", stats.get("max_drawdown", 0)),
        ]
        for row, metrics in enumerate(overall, start=13):
            for col, val in enumerate(metrics, start=1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = Font(name="Aptos", size=10, bold=(col % 2 == 1), color=navy if col % 2 == 1 else "252B25")
                cell.border = thin_border
                if col % 2 == 0 and isinstance(val, (float, int)):
                    metric_label = str(metrics[col - 2]).casefold()
                    if "rate" in metric_label:
                        cell.number_format = percent_format
                    elif metric_label not in ("wins", "losses", "break even", "total trades", "profit factor"):
                        cell.number_format = money_format

        section(19, "CURRENCY PERFORMANCE", 3)
        header_row(20, ["Currency", "Attributed net P&L", "Trades attributed"])
        currency_counts: Dict[str, int] = {}
        for trade in closed_trades:
            normalized = re.sub(r"[^A-Z]", "", str(value(trade, "symbol", "") or "").upper())
            if len(normalized) >= 6 and normalized[:3] in currency_codes and normalized[3:6] in currency_codes:
                currency_counts[normalized[:3]] = currency_counts.get(normalized[:3], 0) + 1
                currency_counts[normalized[3:6]] = currency_counts.get(normalized[3:6], 0) + 1
        currency_rows = sorted(currency_pnl.items(), key=lambda item: item[1], reverse=True)
        best_currency = currency_rows[0][0] if currency_rows else "N/A"
        best_pair = max(pair_pnl, key=pair_pnl.get) if pair_pnl else "N/A"
        ws["E19"] = "TOP PERFORMERS"
        ws["E19"].font = Font(name="Aptos", size=12, bold=True, color=navy)
        ws["E19"].fill = PatternFill("solid", fgColor="E7F3E8")
        ws["E20"] = "Best currency"; ws["F20"] = best_currency
        ws["E21"] = "Best currency P&L"; ws["F21"] = currency_pnl.get(best_currency, 0.0) if best_currency != "N/A" else 0.0
        ws["E22"] = "Best currency pair"; ws["F22"] = best_pair
        ws["E23"] = "Best pair P&L"; ws["F23"] = pair_pnl.get(best_pair, 0.0) if best_pair != "N/A" else 0.0
        ws["E24"] = "Currency attribution"; ws["F24"] = "Base/quote side"
        ws["E25"] = "Method"; ws["F25"] = "BUY: +base / −quote; SELL: reversed"
        opening_capital = float(user_info.get("opening_capital", 0.0) or 0.0)
        ws["E26"] = "Opening capital"; ws["F26"] = opening_capital
        ws["E27"] = "Current balance"; ws["F27"] = opening_capital + float(stats.get("total_net_pnl", 0.0) or 0.0)
        for row in range(20, 28):
            for col in (5, 6):
                ws.cell(row=row, column=col).font = body_font
                ws.cell(row=row, column=col).border = thin_border
                ws.cell(row=row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        ws["F21"].number_format = money_format
        ws["F23"].number_format = money_format
        ws["F26"].number_format = money_format
        ws["F27"].number_format = money_format
        for row, (currency, pnl) in enumerate(currency_rows, start=21):
            ws.cell(row=row, column=1, value=currency).font = body_font
            ws.cell(row=row, column=2, value=pnl).number_format = money_format
            ws.cell(row=row, column=2).font = Font(name="Aptos", size=10, bold=True, color=green if pnl >= 0 else "9A3651")
            ws.cell(row=row, column=3, value=currency_counts.get(currency, 0))
            for col in range(1, 4):
                ws.cell(row=row, column=col).border = thin_border

        if currency_rows:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = "Currency P&L Attribution"
            chart.y_axis.title = "Currency"
            chart.x_axis.title = "Net P&L (USD)"
            chart.height = 6.5
            chart.width = 12
            chart.add_data(Reference(ws, min_col=2, min_row=20, max_row=20 + len(currency_rows)), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=21, max_row=20 + len(currency_rows)))
            ws.add_chart(chart, "H19")

        for col, width in {"A": 22, "B": 19, "C": 17, "D": 14, "E": 22, "F": 38, "G": 3}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A6"

        log = wb.create_sheet("Trade Log")
        log.sheet_view.showGridLines = False
        log.freeze_panes = "A2"
        log_headers = ["Logged at (UTC)", "Close date (UTC)", "Ticket", "Symbol", "Direction", "Lots", "Open price", "Close price", "Stop loss", "Take profit", "Pips", "Gross profit", "Swap", "Commission", "Net P&L", "Position status", "Result"]
        for col, label in enumerate(log_headers, start=1):
            cell = log.cell(row=1, column=col, value=label)
            cell.font = header_font
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        log.row_dimensions[1].height = 28
        for row, trade in enumerate(trades, start=2):
            pnl = float(value(trade, "net_profit", 0.0) or 0.0)
            status = str(value(trade, "status", "CLOSED") or "CLOSED").upper()
            result = "OPEN" if status == "OPEN" else ("WIN" if pnl > 0.001 else "LOSS" if pnl < -0.001 else "BREAK EVEN")
            created = as_utc_naive(value(trade, "created_at"))
            closed_at = as_utc_naive(value(trade, "close_time"))
            row_values = [created, closed_at, value(trade, "ticket") or "—", str(value(trade, "symbol", "") or "").upper(), str(value(trade, "trade_type", "") or "").upper(), value(trade, "volume", 0), value(trade, "open_price", 0), value(trade, "close_price", 0), value(trade, "sl"), value(trade, "tp"), value(trade, "pips", 0), value(trade, "profit", 0), value(trade, "swap", 0), value(trade, "commission", 0), pnl, status, result]
            for col, val in enumerate(row_values, start=1):
                cell = log.cell(row=row, column=col, value=val)
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left" if col in (3, 4, 17) else "center")
                cell.border = thin_border
                if row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F7F9F6")
            for col in (1, 2):
                log.cell(row=row, column=col).number_format = "yyyy-mm-dd hh:mm"
            for col in (7, 8, 9, 10):
                log.cell(row=row, column=col).number_format = "0.00000"
            log.cell(row=row, column=6).number_format = "0.00"
            log.cell(row=row, column=11).number_format = "0.0"
            for col in (12, 13, 14, 15):
                log.cell(row=row, column=col).number_format = money_format
            pnl_fill = pale_green if pnl > 0 else pale_red if pnl < 0 else None
            if pnl_fill:
                log.cell(row=row, column=15).fill = PatternFill("solid", fgColor=pnl_fill)
                log.cell(row=row, column=17).fill = PatternFill("solid", fgColor=pnl_fill)
        log.auto_filter.ref = f"A1:Q{max(1, len(trades) + 1)}"
        for col, width in enumerate([20, 20, 15, 14, 12, 10, 14, 14, 14, 14, 10, 16, 12, 14, 16, 16, 14], start=1):
            log.column_dimensions[get_column_letter(col)].width = width
        log.sheet_properties.pageSetUpPr.fitToPage = True

        wb.properties.title = "MyFXkeeper Trading Performance Report"
        wb.properties.subject = "WTD, MTD, YTD performance and trade history"
        wb.properties.creator = "MyFXkeeper"
        wb.save(output_path)
        return True
    except Exception:
        logger.exception("Excel report generation failed")
        return False
